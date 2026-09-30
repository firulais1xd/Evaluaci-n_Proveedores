"""Lectura del archivo EvaluacionResumenPeriodo, cálculo de resultados, comparación y textos."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import load_workbook

CATS = ["EXCELENTE", "MUY BUENO", "BUENO", "REGULAR", "MALO"]
CAT_LABEL = {"EXCELENTE": "Excelente", "MUY BUENO": "Muy bueno", "BUENO": "Bueno", "REGULAR": "Regular", "MALO": "Malo"}
CAT_PLURAL = {"EXCELENTE": ("excelente", "excelentes"), "MUY BUENO": ("muy bueno", "muy buenos"),
              "BUENO": ("bueno", "buenos"), "REGULAR": ("regular", "regulares"), "MALO": ("malo", "malos")}
CAT_HEX = {"EXCELENTE": "#11774A", "MUY BUENO": "#4FA877", "BUENO": "#CFA11D", "REGULAR": "#D65F28", "MALO": "#A3233A"}
# Umbral mínimo para pasar a la siguiente categoría
NEXT = {"MUY BUENO": (98, "excelente"), "BUENO": (80, "muy bueno"), "REGULAR": (70, "bueno"), "MALO": (60, "regular")}


def cat_of(v: float) -> str:
    if v >= 98:
        return "EXCELENTE"
    if v >= 80:
        return "MUY BUENO"
    if v >= 70:
        return "BUENO"
    if v >= 60:
        return "REGULAR"
    return "MALO"


def cat_rank(c: str) -> int:
    return CATS.index(c)


PROC = {"CP": "Compras nacionales", "CE": "Comercio exterior", "MN": "Mantenimiento", "GH": "Gestión humana",
        "ST": "Sistemas", "GA": "Gestión ambiental", "VT": "Ventas", "CC": "Control de calidad", "MT": "Metrología",
        "GI": "Gestión de la inocuidad", "SB": "Seguridad física", "GC": "Gestión de calidad", "EN": "Preparación",
        "OT": "Otros"}
PROC_ORDER = list(PROC)

# Nombres cortos por código de evaluación. Si aparece un código nuevo, el nombre se toma del título de la hoja.
TYPE_NAMES = {
    "CP-DS-16": "Empaques y embalajes", "CP-DS-01": "Insumos y productos químicos", "CP-DS-02": "Repuestos y equipos",
    "CP-DS-19": "Aparatos electrónicos y eléctricos", "CP-DS-20": "Llantas",
    "CE-DS-02": "Proveedores internacionales de químicos",
    "CE-DS-03": "Proveedores internacionales de repuestos, equipos y accesorios",
    "CE-DS-01": "Proveedores internacionales de gelatina", "CE-DS-06": "Proveedores internacionales de materia prima",
    "CE-DS-05": "Transporte marítimo", "CE-RG-03": "Transporte terrestre para Comex", "CE-RG-02": "Otros servicios Comex",
    "MN-DS-03": "Mantenimiento de vehículos y/o equipos", "MN-RG-105": "Mantenimiento de equipos de refrigeración",
    "MN-DS-05": "Construcción de obras civiles",
    "GH-DS-03-01": "Estudios de confiabilidad (poligrafías)", "GH-RG-60": "Transporte de personal",
    "GH-DS-05-1": "Servicios temporales", "GH-DS-08": "Carné e identificación de personal",
    "GH-RG-64": "Exámenes médicos ocupacionales",
    "ST-DS-02": "Servicios informáticos (internet y equipos)", "ST-DS-01": "Servicios informáticos (programación)",
    "GA-RG-39": "Asesoría y consultoría ambiental", "GA-RG-10": "Monitoreos ambientales",
    "GA-RG-12": "Disposición de residuos ordinarios", "GA-RG-08": "Gestores de residuos",
    "GA-RG-38": "Mantenimiento de jardinería",
    "VT-RG-25": "Transporte de producto terminado", "2024": "Servicios de laboratorio subcontratado",
    "MT-RG-43": "Servicios metrológicos", "GI-RG-04": "Manejo integral de plagas",
    "SB-RG-28": "Seguridad física (escoltas)", "SB-DS-07": "Seguridad satelital (apertura de puertas de contenedor)",
    "SB-RG-30": "Mantenimiento de CCTV", "GC-DS-31": "Entes certificadores", "EN-RG-36": "Carnaza nacional",
}


# Título de cada evaluación (sin el código). Sirve para reconocer la hoja cuando la plataforma exporta el
# código incompleto (p. ej. "EV-12-02" en lugar de "EV-12-CE-RG-02") o lo cambia entre años.
KNOWN_TITLES = {
    "CE-DS-03": "EVALUACIÓN DE DESEMPEÑO PROVEEDORES INTERNACIONALES DE REPUESTOS, EQUIPOS Y ACCESORIOS",
    "MN-DS-03": "EVALUACION PROVEEDORES DE SERVICIO DE MTTO VEHICULOS Y/O EQUIPOS",
    "CP-DS-02": "EVALUACIÓN DE DESEMPEÑO DE PROVEEDORES NACIONALES REPUESTOS Y EQUIPOS",
    "CE-DS-02": "EVALUACIÓN DE DESEMPEÑO PROVEEDORES INTERNACIONALES DE QUÍMICOS",
    "CE-DS-01": "EVALUACIÓN DE DESEMPEÑO PROVEEDORES INTERNACIONALES DE GELATINA",
    "GH-DS-03-01": "EVALUACIÓN DE DESEMPEÑO PROVEEDOR DE POLIGRAFIAS Y ESTUDIOS DE CONFIABILIDAD AL PERSONAL",
    "ST-DS-02": "EVALUACIÓN DE DESEMPEÑO DE PROVEEDOR SERVICIOS INFORMÁTICOS (INTERNET YEQUIPOS)",
    "MN-RG-105": "EVALUACION PROVEEDOR DE MTTO DE EQUIPOS DE REFRIGERACION.",
    "CP-DS-16": "EVALUACIÓN DE DESEMPEÑO DE PROVEEDORES DE EMPAQUES Y EMBALAJES",
    "CP-DS-01": "EVALUACIÓN DE DESEMPEÑO DE PROVEEDOR DE INSUMOS Y PRODUCTOS QUÍMICOS",
    "CP-DS-19": "APARATOS ELECTRONICOS Y ELECTRICOS",
    "CE-RG-02": "EVALUACIÓN DE DESEMPEÑO PROVEEDORES DE SERVICIOS PARA COMEX",
    "GA-RG-12": "EVALUACIÓN DE DESEMPEÑO PROVEEDOR DE SERVICIO DE DISPOSICIÓN DE RESIDUOS ORDINARIOS",
    "GA-RG-08": "EVALUACIÓN DE DESEMPEÑO DE GESTORES DE RESIDUOS",
    "MN-DS-05": "EVALUACION PROVEEDORES DE SERVICIO DE CONSTRUCCION OBRAS CIVILES",
    "CE-DS-05": "EVALUACIÓN DE DESEMPEÑO PROVEEDORES DE TRANSPORTE MARÍTIMO",
    "GA-RG-10": "EVALUACIÓN DE DESEMPEÑO DE PROVEEDOR DE SERVICIO DE MONITOREOS AMBIENTALES",
    "2024": "EVALUACIÓN DE DESEMPEÑO DE PROVEEDOR DE SERVICIO SUBCONTRATADO DE LABORATORIO",
    "CE-RG-03": "EVALUACIÓN DE DESEMPEÑO PROVEEDORES DE TRANSPORTE TERRESTRE PARA COMEX",
    "MT-RG-43": "EVALUACION DE PROVEEDORES DE SERVICIOS METROLOGICOS",
    "GH-DS-05-1": "EVALUACIÓN DE DESEMPEÑO EMPRESA DE SERVICIOS TEMPORALES.",
    "GI-RG-04": "EVALUACIÓN DE DESEMPEÑO PROVEEDOR DE CONTROL DE PLAGAS",
    "CP-DS-20": "EVALUACIÓN DE PROVEEDORES DE LLANTAS",
    "GC-DS-31": "EVALUACIÓN A ENTES CERTIFICADORES",
    "SB-RG-28": "EVALUACIÓN DE DESEMPEÑO DE PROVEEDOR DE SERVICIOS DE SEGURIDAD FÍSICA.",
    "GA-RG-39": "EVALUACIÓN DE PROVEEDOR DE SERVICIO DE ASESORIA Y CONSULTORÍA AMBIENTAL",
    "VT-RG-25": "EVALUACIÓN DE DESEMPEÑO PROVEEDORES DE TRANSPORTE TERRESTRE (VENTAS)",
    "EN-RG-36": "EVALUACION DE DESEMPEÑO DE PROVEEDORES CARNAZA NACIONAL",
    "GH-RG-60": "EVALUACIÓN DE DESEMPEÑO PROVEEDORES DE SERVICIO DE TRANSPORTE.",
    "GA-RG-38": "EVALUACIÓN PROVEEDOR SERVICIOS DE MANTENIMIENTO DE JARDINERIA",
    "ST-DS-01": "EVALUACION DE DESEMPEÑO DE PROVEEDOR SERVICIOS INFORMATICOS (PROGRAMACION)",
    "SB-DS-07": "EVALUACION DE DESEMPEÑO PROVEEDOR SEGURIDAD SATELITAL APERTURA PUERTAS DEL CONTENEDOR",
    "GH-DS-08": "IDENTIFICACIÓN DE PERSONAL",
    "GH-RG-64": "EVALUACIÓN DE DESEMPEÑO DE PROVEEDOR DE EXÁMENES MÉDICOS OCUPACIONALES",
    "SB-RG-30": "EVALUACION DE DESEMPEÑO DE PROVEEDOR MANTENIMIENTO CCTV",
    "CE-DS-06": "EVALUACIÓN DE DESEMPEÑO PROVEEDORES INTERNACIONALES DE MATERIA PRIMA",
}


# ---------------------------------------------------------------- utilidades
def clean(s) -> str:
    return re.sub(r"\s+", " ", "" if s is None else str(s)).strip()


def norm_nit(s) -> str:
    return re.sub(r"\s", "", clean(s)).upper()


def to_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace("%", "").replace(",", "."))
    except ValueError:
        return None


def fnum(v: float, dec: int = 2) -> str:
    s = f"{round(v, dec):.{dec}f}".rstrip("0").rstrip(".")
    return s.replace(".", ",") if s not in ("-0",) else "0"


def pct(v) -> str:
    return "—" if v is None else fnum(v) + "%"


def pp(d: float) -> str:
    sign = "+" if d > 0.005 else "−" if d < -0.005 else "±"
    return f"{sign}{fnum(abs(d))} pp"


def share(a, b) -> float:
    return round(a / b * 1000) / 10 if b else 0.0


def join_es(items) -> str:
    items = list(items)
    if len(items) <= 1:
        return items[0] if items else ""
    return ", ".join(items[:-1]) + " y " + items[-1]


def title_no_code(t: str) -> str:
    t = clean(t)
    i = t.find(" - ")
    return t[i + 3:].strip() if 0 <= i < 24 else t


def type_name_from(code: str, title: str) -> str:
    if code in TYPE_NAMES:
        return TYPE_NAMES[code]
    t = title_no_code(title).rstrip(".")
    t = re.sub(r"^EVALUACI[OÓ]N\s+(DE\s+)?(DESEMPE[NÑ]O\s+)?(DE\s+|A\s+)?(PROVEEDOR(ES)?\s+)?(DE\s+)?", "", t,
               flags=re.I)
    t = (t or code).lower()
    return t[:1].upper() + t[1:]


def _norm_title(t: str) -> str:
    import unicodedata
    t = unicodedata.normalize("NFKD", title_no_code(t or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z0-9]+", "", t)


_TITLE_TO_CODE = {_norm_title(v): k for k, v in KNOWN_TITLES.items()}


def canon_code(code: str, title: str) -> str:
    """Código estable de la evaluación: si el código no es uno conocido, se busca por el título."""
    if code in TYPE_NAMES:
        return code
    return _TITLE_TO_CODE.get(_norm_title(title), code)


def proc_of(code: str, title: str) -> str:
    m = re.match(r"^([A-Z]{2})-", code)
    if m and m.group(1) in PROC:
        return m.group(1)
    if re.search("LABORATORIO", title, re.I):
        return "CC"
    if re.search(r"COMEX|INTERNACIONAL", title, re.I):
        return "CE"
    return "OT"


def pkey(r: dict) -> str:
    return norm_nit(r["nit"]) or r["name"].upper()


# ---------------------------------------------------------------- lectura
def parse_workbook(data: bytes, label: str) -> dict:
    """Lee el archivo de la plataforma: una hoja por tipo de evaluación."""
    wb = load_workbook(io.BytesIO(data), data_only=True)  # modo normal: algunos archivos declaran mal su tamaño
    sheets = []
    for ws in wb.worksheets:
        aoa = [list(r) for r in ws.iter_rows(values_only=True)]
        hi = next((i for i, r in enumerate(aoa) if any(clean(c).lower() == "nitproveedor" for c in r)), -1)
        if hi < 0:
            continue
        head = [clean(c) for c in aoa[hi]]
        title = ""
        for r in aoa[:hi]:
            c = next((clean(x) for x in r if clean(x)), "")
            if c:
                title = c
                break

        def idx(k):
            return next((i for i, h in enumerate(head) if h.lower() == k), -1)

        i_nit, i_name, i_fin, i_cat = idx("nitproveedor"), idx("nombreproveedor"), idx("final"), idx("categoria")
        if i_fin < 0:
            continue
        crit = [(h, i) for i, h in enumerate(head) if h and i not in (i_nit, i_name, i_fin, i_cat)]
        rows, skipped = [], []
        for r in aoa[hi + 1:]:
            r = r + [None] * (len(head) - len(r))
            nit, nm = clean(r[i_nit]), clean(r[i_name])
            if not nit and not nm:
                continue
            fin = to_num(r[i_fin])
            if fin is None:
                skipped.append(nm or nit)
                continue
            plat = clean(r[i_cat]).upper() if i_cat >= 0 else ""
            rows.append({"nit": nit, "name": nm, "final": fin, "platCat": plat,
                         "scores": [to_num(r[i]) for _, i in crit]})
        used = [j for j in range(len(crit)) if any(x["scores"][j] is not None for x in rows)]
        for x in rows:
            x["scores"] = [x["scores"][j] for j in used]
        code = canon_code(re.sub(r"^EV-\d+-", "", ws.title, flags=re.I), title)
        sheets.append({"sheet": ws.title, "code": code, "title": title, "criteria": [crit[j][0] for j in used],
                       "rows": rows, "skipped": skipped})
    wb.close()
    if not sheets:
        raise ValueError("No encontré hojas con la columna NitProveedor. ¿Es el archivo EvaluacionResumenPeriodo?")
    return {"label": label, "savedAt": None, "sheets": sheets}


def to_stored(p: dict) -> dict:
    r4 = lambda v: None if v is None else round(v, 4)
    return {"label": p["label"], "savedAt": datetime.now().isoformat(timespec="seconds"),
            "sheets": [{"sheet": s["sheet"], "code": s["code"], "title": s["title"], "criteria": s["criteria"],
                        "skipped": s.get("skipped", []),
                        "rows": [[r["nit"], r["name"], r4(r["final"]), r["platCat"], [r4(v) for v in r["scores"]]]
                                 for r in s["rows"]]} for s in p["sheets"]]}


def from_stored(d: dict) -> dict:
    return {"label": d["label"], "savedAt": d.get("savedAt"),
            "sheets": [{**s, "code": canon_code(s["code"], s.get("title", "")), "skipped": s.get("skipped", []),
                        "rows": [{"nit": a[0], "name": a[1], "final": a[2], "platCat": a[3] or "",
                                  "scores": a[4] or []} for a in s["rows"]]} for s in d.get("sheets", [])]}


# ---------------------------------------------------------------- recálculo del Final
# La plataforma calcula el Final como promedio PONDERADO de los criterios y toma los criterios
# en blanco como 0. Aquí se deducen los pesos de cada tipo de evaluación a partir de los propios
# datos (Final = suma de peso x puntaje) y, en las filas con criterios en blanco, se recalcula el
# Final sin tenerlos en cuenta:  Final nuevo = Final plataforma / (1 - peso de los criterios en blanco).
# Las filas completas conservan el Final de la plataforma.
def infer_weights(periods: list[dict]) -> dict:
    import numpy as np

    pool: dict = {}
    for p in periods:
        for s in p["sheets"]:
            if s["criteria"]:
                pool.setdefault((s["code"], tuple(s["criteria"])), []).extend(s["rows"])
    out = {}
    for (code, crit), rows in pool.items():
        k = len(crit)
        rows = [r for r in rows if len(r["scores"]) == k]
        if not rows:
            continue
        X = np.array([[0.0 if v is None else v / 100 for v in r["scores"]] for r in rows])
        y = np.array([r["final"] / 100 for r in rows])
        A = np.vstack([X, np.ones(k)])        # los pesos suman 100%
        b = np.append(y, 1.0)
        w, *_ = np.linalg.lstsq(A, b, rcond=None)
        rank = np.linalg.matrix_rank(A)
        exact = bool(np.abs(A @ w - b).max() < 5e-4 and (w > -1e-3).all())
        ident = [bool(np.linalg.matrix_rank(np.vstack([A, np.eye(k)[j]])) == rank) for j in range(k)]
        out[(code, crit)] = {"code": code, "criteria": list(crit), "weights": [float(x) for x in w],
                             "exact": exact, "ident": ident, "n": len(rows)}
    return out


def recalc_period(p: dict, wmap: dict, overrides: dict | None = None, ignore_blanks: bool = True) -> dict:
    """Devuelve una copia del periodo con el Final recalculado en las filas con criterios en blanco."""
    overrides = overrides or {}
    sheets = []
    for s in p["sheets"]:
        crit = s["criteria"]
        k = len(crit)
        info = wmap.get((s["code"], tuple(crit)))
        ov = overrides.get(s["code"])
        ov_w = [ov.get(c) for c in crit] if isinstance(ov, dict) else None
        if ov_w and (None in ov_w or sum(ov_w) <= 0):
            ov_w = None
        rows = []
        for r in s["rows"]:
            r2 = {**r, "finalPlat": r["final"], "blanks": [], "recalc": False, "estimated": False}
            blanks = [j for j, v in enumerate(r["scores"]) if v is None] if len(r["scores"]) == k else []
            if ignore_blanks and blanks and len(blanks) < k:
                r2["blanks"] = [crit[j] for j in blanks]
                if ov_w:
                    tot = sum(ov_w)
                    wb = sum(ov_w[j] for j in blanks) / tot
                    est = False
                elif info and info["exact"] and all(info["ident"][j] for j in blanks):
                    wb = sum(info["weights"][j] for j in blanks)
                    est = False
                else:  # peso desconocido: se supone igual para todos los criterios
                    wb = len(blanks) / k
                    est = True
                if 0 < wb < 0.999:
                    r2["final"] = min(100.0, r["final"] / (1 - wb))
                    r2["recalc"] = abs(r2["final"] - r["final"]) > 1e-6
                    r2["estimated"] = est
            rows.append(r2)
        sheets.append({**s, "rows": rows})
    return {**p, "sheets": sheets}


def recalc_rows(a: dict) -> list:
    """Filas cuyo Final cambió por criterios en blanco: [(fila, tipo)]."""
    return [(r, t) for t in a["types"] for r in t.rows if r.get("recalc")]


# ---------------------------------------------------------------- análisis
@dataclass
class Tipo:
    code: str
    sheet: str
    proc: str
    title: str
    type_name: str
    rows: list
    n: int
    counts: dict
    avg: float
    min: float
    max: float
    crit: list
    weakest: dict | None
    skipped: list = field(default_factory=list)


def _counts():
    return {c: 0 for c in CATS}


def analyze(p: dict) -> dict:
    types = []
    for s in p["sheets"]:
        if not s["rows"]:
            continue
        rows = sorted(({**r, "cat": cat_of(r["final"])} for r in s["rows"]), key=lambda r: (-r["final"], r["name"]))
        n = len(rows)
        counts = _counts()
        for r in rows:
            counts[r["cat"]] += 1
        crit = []
        for j, h in enumerate(s["criteria"]):
            v = [r["scores"][j] for r in rows if j < len(r["scores"]) and r["scores"][j] is not None]
            if v:
                crit.append({"h": h, "j": j, "avg": sum(v) / len(v)})
        weakest = min(crit, key=lambda c: c["avg"]) if len(crit) > 1 else None
        types.append(Tipo(code=s["code"], sheet=s["sheet"], proc=proc_of(s["code"], s["title"]),
                          title=title_no_code(s["title"]), type_name=type_name_from(s["code"], s["title"]),
                          rows=rows, n=n, counts=counts, avg=sum(r["final"] for r in rows) / n,
                          min=rows[-1]["final"], max=rows[0]["final"], crit=crit, weakest=weakest,
                          skipped=s.get("skipped", [])))
    all_rows = [(r, t) for t in types for r in t.rows]
    counts = _counts()
    for r, _ in all_rows:
        counts[r["cat"]] += 1
    procs = []
    for k in PROC_ORDER:
        ts = [t for t in types if t.proc == k]
        if not ts:
            continue
        rs = [r for t in ts for r in t.rows]
        c = _counts()
        for r in rs:
            c[r["cat"]] += 1
        procs.append({"k": k, "name": PROC[k], "types": ts, "n": len(rs), "counts": c,
                      "avg": sum(r["final"] for r in rs) / len(rs), "low": sum(r["final"] < 80 for r in rs)})
    alerts = []
    for t in types:
        for r in t.rows:
            if not r["platCat"]:
                alerts.append({"t": t, "r": r, "kind": "sin"})
            elif r["platCat"] != cat_of(r.get("finalPlat", r["final"])):
                alerts.append({"t": t, "r": r, "kind": "dif"})
        for nm in t.skipped:
            alerts.append({"t": t, "r": {"name": nm}, "kind": "nofinal"})
    total = len(all_rows)
    return {"label": p["label"], "types": types, "all": all_rows, "counts": counts, "total": total,
            "uniq": len({pkey(r) for r, _ in all_rows}),
            "avg": sum(r["final"] for r, _ in all_rows) / total if total else None,
            "procs": procs, "alerts": alerts}


def compare(cur: dict, prev: dict | None) -> dict | None:
    if not prev:
        return None
    pt = {t.code: t for t in prev["types"]}
    by_type, movers = {}, []
    tot = dict(improved=0, declined=0, same=0, matched=0, nNew=0, catUp=0, catDown=0)
    for t in cur["types"]:
        p = pt.get(t.code)
        res = {"prev": p, "matched": [], "news": [], "gone": [], "improved": 0, "declined": 0, "same": 0}
        if p:
            pm = {pkey(r): r for r in p.rows}
            seen = set()
            for r in t.rows:
                k = pkey(r)
                pr = pm.get(k)
                if pr:
                    seen.add(k)
                    d = r["final"] - pr["final"]
                    m = {"r": r, "pr": pr, "d": d, "t": t}
                    res["matched"].append(m)
                    movers.append(m)
                    if d > 0.05:
                        res["improved"] += 1
                    elif d < -0.05:
                        res["declined"] += 1
                    else:
                        res["same"] += 1
                    dr = cat_rank(pr["cat"]) - cat_rank(r["cat"])
                    tot["catUp"] += dr > 0
                    tot["catDown"] += dr < 0
                else:
                    res["news"].append(r)
            res["gone"] = [r for r in p.rows if pkey(r) not in seen]
        else:
            res["news"] = list(t.rows)
        for k in ("improved", "declined", "same"):
            tot[k] += res[k]
        tot["matched"] += len(res["matched"])
        tot["nNew"] += len(res["news"])
        by_type[t.code] = res
    cur_codes = {t.code for t in cur["types"]}
    return {"prevLabel": prev["label"], "prev": prev, "byType": by_type, **tot,
            "goneTypes": [t for t in prev["types"] if t.code not in cur_codes],
            "movers": sorted(movers, key=lambda m: m["d"])}


# ---------------------------------------------------------------- textos
def _names_at(rows, v):
    ns = [r["name"] for r in rows if abs(r["final"] - v) < 1e-9]
    return join_es(ns) if len(ns) <= 2 else f"{ns[0]} y {len(ns) - 1} proveedores más"


def vb(n: int, sing: str, plur: str) -> str:
    return f"{n} {sing if n == 1 else plur}"


def counts_text(c: dict) -> str:
    return join_es([f"{c[k]} {CAT_PLURAL[k][0 if c[k] == 1 else 1]}" for k in CATS if c[k]])


def type_paras(t: Tipo, cmp: dict | None) -> list[str]:
    out = []
    if t.n == 1:
        out.append(f"En esta categoría se cuenta con 1 proveedor, que obtuvo una calificación de {pct(t.max)} "
                   f"– {t.rows[0]['name']} –.")
    elif abs(t.max - t.min) < 1e-9:
        out.append(f"En esta categoría se cuenta con {t.n} proveedores; todos obtuvieron una calificación de {pct(t.max)}.")
    else:
        out.append(f"En esta categoría se cuenta con {t.n} proveedores, que obtuvieron calificaciones que oscilan entre "
                   f"{pct(t.min)} – {_names_at(t.rows, t.min)} – y {pct(t.max)} – {_names_at(t.rows, t.max)} –.")
    if t.n > 1:
        out.append(f"Promedio de la categoría: {pct(t.avg)}. Distribución: {counts_text(t.counts)}.")
    low = sorted((r for r in t.rows if r["final"] < 80), key=lambda r: r["final"])
    if low:
        shown = [f"{r['name']} ({pct(r['final'])}, {CAT_PLURAL[r['cat']][0]})" for r in low[:6]]
        if len(low) > 6:
            shown.append(f"{len(low) - 6} más")
        out.append(f"Requieren seguimiento por calificación inferior a 80%: {join_es(shown)}.")
    if t.weakest and t.weakest["avg"] < 99.95:
        out.append(f"El criterio con menor cumplimiento fue «{t.weakest['h']}», con un promedio de {pct(t.weakest['avg'])}.")
    if cmp:
        c = cmp["byType"].get(t.code)
        if not c or not c["prev"]:
            out.append(f"Este tipo de evaluación no tuvo registros en {cmp['prevLabel']}.")
        else:
            pv = c["prev"]
            d = t.avg - pv.avg
            if t.n == 1 and pv.n == 1 and len(c["matched"]) == 1:
                s = f"Frente a {cmp['prevLabel']}, su calificación pasó de {pct(pv.avg)} a {pct(t.avg)} ({pp(d)})."
            else:
                s = f"Frente a {cmp['prevLabel']}, el promedio pasó de {pct(pv.avg)} a {pct(t.avg)} ({pp(d)})."
                if len(c["matched"]) == 1:
                    s += (" El único proveedor evaluado en ambos periodos "
                          + ("mejoró" if c["improved"] else "bajó" if c["declined"] else "se mantuvo") + ".")
                elif c["matched"]:
                    s += (f" De los {len(c['matched'])} proveedores evaluados en ambos periodos, "
                          f"{vb(c['improved'], 'mejoró', 'mejoraron')}, {vb(c['declined'], 'bajó', 'bajaron')} y "
                          f"{vb(c['same'], 'se mantuvo', 'se mantuvieron')}.")
            if c["news"]:
                s += " Ingresó 1 proveedor nuevo." if len(c["news"]) == 1 else f" Ingresaron {len(c['news'])} proveedores nuevos."
            if c["gone"]:
                s += (" 1 proveedor del periodo anterior no fue evaluado en este periodo." if len(c["gone"]) == 1
                      else f" {len(c['gone'])} proveedores del periodo anterior no fueron evaluados en este periodo.")
            worst = min(c["matched"], key=lambda m: m["d"], default=None)
            if t.n > 1 and worst and worst["d"] <= -5:
                s += (f" La mayor caída fue la de {worst['r']['name']} (de {pct(worst['pr']['final'])} a "
                      f"{pct(worst['r']['final'])}).")
            out.append(s)
    return out


def consolidated_paras(a: dict, cmp: dict | None) -> list[str]:
    c, T = a["counts"], a["total"]
    out = [f"Se verificaron {T} evaluaciones, correspondientes a {a['uniq']} proveedores (un mismo proveedor puede "
           f"evaluarse en más de un tipo). El promedio general fue {pct(a['avg'])}.",
           "De las evaluaciones verificadas, la mayoría de los proveedores se ubica en las categorías "
           + join_es([f"“{k}” ({fnum(share(c[k], T))}%)" for k in CATS[:2]]) + "."]
    low_cats = [k for k in CATS[2:] if c[k]]
    if low_cats:
        lows = [(r, t) for r, t in a["all"] if r["final"] < 80]
        procs = list(dict.fromkeys(PROC[t.proc].lower() for _, t in lows))
        parts = [f"{c[k]} se {'ubicó' if c[k] == 1 else 'ubicaron'} en la categoría “{k}” ({fnum(share(c[k], T))}%)"
                 for k in low_cats]
        out.append(f"Del total de proveedores evaluados, {join_es(parts)}. "
                   f"{'Este proveedor pertenece' if len(lows) == 1 else 'Estos proveedores pertenecen'} "
                   f"{'al proceso' if len(procs) == 1 else 'a los procesos'} de {join_es(procs)}.")
    else:
        out.append("Ningún proveedor quedó por debajo de 80%.")
    ranked = sorted(a["procs"], key=lambda g: g["avg"])
    if len(ranked) > 1:
        out.append(f"El proceso con menor promedio fue {ranked[0]['name'].lower()} ({pct(ranked[0]['avg'])}) y el de "
                   f"mayor promedio, {ranked[-1]['name'].lower()} ({pct(ranked[-1]['avg'])}).")
    if cmp:
        pa = cmp["prev"]
        out.append(f"Frente a {cmp['prevLabel']} ({pa['total']} evaluaciones), el promedio general pasó de "
                   f"{pct(pa['avg'])} a {pct(a['avg'])} ({pp(a['avg'] - pa['avg'])}) y la proporción de proveedores "
                   f"excelentes pasó de {fnum(share(pa['counts']['EXCELENTE'], pa['total']))}% a "
                   f"{fnum(share(c['EXCELENTE'], T))}%. De los {cmp['matched']} proveedores evaluados en ambos periodos, "
                   f"{vb(cmp['improved'], 'mejoró', 'mejoraron')} su calificación, {vb(cmp['declined'], 'bajó', 'bajaron')} y "
                   f"{vb(cmp['same'], 'se mantuvo', 'se mantuvieron')}; {vb(cmp['catUp'], 'subió', 'subieron')} de categoría "
                   f"y {vb(cmp['catDown'], 'bajó', 'bajaron')} de categoría.")
    return out


def conclusion_text(a: dict, cmp: dict | None) -> str:
    T = a["total"]
    good = a["counts"]["EXCELENTE"] + a["counts"]["MUY BUENO"]
    low = sum(r["final"] < 80 for r, _ in a["all"])
    s = (f"En términos generales, el desempeño global de proveedores es "
         f"{'satisfactorio' if share(good, T) >= 90 else 'aceptable, con oportunidades de mejora'}: el "
         f"{fnum(share(good, T))}% de las evaluaciones se ubicó en las categorías Excelente o Muy bueno.")
    if low:
        s += f" Se recomienda hacer seguimiento a {'el proveedor' if low == 1 else f'los {low} proveedores'} con calificación inferior a 80%."
    else:
        s += " Ningún proveedor obtuvo calificación inferior a 80%."
    if cmp:
        d = a["avg"] - cmp["prev"]["avg"]
        est = "se mantuvo estable" if abs(d) < 0.05 else "mejoró" if d > 0 else "disminuyó"
        s += f" Frente a {cmp['prevLabel']}, el promedio general {est} ({pp(d)})."
    s += "\n\nSe evidenció que las evaluaciones fueron gestionadas oportunamente y durante los tiempos establecidos."
    return s


def action_rows(a: dict) -> list[list[str]]:
    rows = []
    for r, t in sorted(((r, t) for r, t in a["all"] if r["final"] < 80), key=lambda x: x[0]["final"]):
        if r["final"] < 70:
            rows.append(["AC", f"Solicitar plan de acción a {r['name']} ({t.type_name}, {pct(r['final'])}, "
                               f"{CAT_LABEL[r['cat']].lower()})."])
        else:
            rows.append(["CR", f"Seguimiento al desempeño de {r['name']} ({t.type_name}, {pct(r['final'])}) en la "
                               f"próxima evaluación."])
    return rows or [["NA", "NA"]]


# ---------------------------------------------------------------- proveedores
def provider_index(a: dict) -> dict:
    idx = {}
    for t in a["types"]:
        for r in t.rows:
            k = pkey(r)
            idx.setdefault(k, {"key": k, "nit": r["nit"], "name": r["name"], "evals": []})["evals"].append((t, r))
    return idx


def rank_in(t: Tipo, r: dict) -> int:
    return sum(x["final"] > r["final"] + 1e-9 for x in t.rows) + 1


def crit_rows(t: Tipo, r: dict) -> list[dict]:
    out = []
    for c in t.crit:
        v = r["scores"][c["j"]] if c["j"] < len(r["scores"]) else None
        if v is not None:
            out.append({"h": c["h"], "v": v, "avg": c["avg"]})
    return out


def history_for(timeline: list[tuple[str, dict]], key: str, code: str) -> list[dict]:
    out = []
    for label, a in timeline:
        t = next((t for t in a["types"] if t.code == code), None)
        r = next((r for r in t.rows if pkey(r) == key), None) if t else None
        if r:
            out.append({"label": label, "final": r["final"], "cat": r["cat"]})
    return out


def provider_insights(t: Tipo, r: dict, others: list, hist: list, cmp: dict | None) -> list[str]:
    out = []
    s = f"Calificación final de {pct(r['final'])}, categoría {CAT_LABEL[r['cat']].lower()}."
    if t.n > 1:
        diff = r["final"] - t.avg
        pos = ("en el promedio del tipo" if abs(diff) < 0.005 else
               f"{fnum(abs(diff))} pp {'por encima' if diff > 0 else 'por debajo'} del promedio del tipo")
        s += f" Ocupa el puesto {rank_in(t, r)} de {t.n} en {t.type_name.lower()} y está {pos} ({pct(t.avg)})."
    out.append(s)
    if r["cat"] in NEXT:
        th, lab = NEXT[r["cat"]]
        out.append(f"Le faltan {fnum(th - r['final'])} pp para pasar a la categoría {lab} ({th}%).")
    if r.get("recalc"):
        out.append(f"La plataforma reporta {pct(r['finalPlat'])} porque toma como 0 "
                   + join_es(["«" + b + "»" for b in r["blanks"]])
                   + f", que quedó en blanco; sin {'ese criterio' if len(r['blanks']) == 1 else 'esos criterios'} la "
                   f"calificación es {pct(r['final'])}" + (" (estimada: se supuso el mismo peso para todos los criterios)"
                                                           if r.get("estimated") else "") + ".")
    cr = crit_rows(t, r)
    weak = sorted((c for c in cr if c["v"] < 100), key=lambda c: c["v"])
    if cr and not weak:
        out.append(f"Cumplió al 100% en los {len(cr)} criterios evaluados.")
    if cr and r["final"] < min(c["v"] for c in cr) - 0.5:
        out.append(f"Atención: la calificación final ({pct(r['final'])}) es menor que todos sus criterios (el más bajo "
                   f"es {pct(min(c['v'] for c in cr))}); conviene revisar en la plataforma si falta un criterio o una "
                   f"ponderación.")
    if weak:
        out.append("Por mejorar: " + join_es([f"«{c['h']}» con {pct(c['v'])}"
                                              + (f" (promedio del tipo {pct(c['avg'])})" if t.n > 1 else "")
                                              for c in weak[:3]]) + ".")
        zero = [c for c in weak if c["v"] == 0]
        if zero:
            out.append(f"Obtuvo 0% en {join_es(['«' + c['h'] + '»' for c in zero])}; conviene verificar si el criterio no "
                       f"aplicaba o si hubo un incumplimiento total, porque afecta la calificación final.")
        strong = sorted((c for c in cr if c["v"] >= 99.95 and c["avg"] < 95), key=lambda c: c["avg"])[:2]
        if strong:
            out.append("Se destaca en " + join_es([f"«{c['h']}» (100% frente a {pct(c['avg'])} del tipo)"
                                                   for c in strong]) + ".")
    if len(hist) >= 2:
        f0, f1 = hist[0], hist[-1]
        extra = (f"; pasó de {CAT_LABEL[f0['cat']].lower()} a {CAT_LABEL[f1['cat']].lower()}"
                 if f0["cat"] != f1["cat"] else "")
        out.append("Evolución: " + " → ".join(f"{h['label']} {pct(h['final'])}" for h in hist)
                   + f" ({pp(f1['final'] - f0['final'])}{extra}).")
    elif cmp:
        out.append(f"No aparece en {t.type_name.lower()} en {cmp['prevLabel']}; es su primera evaluación registrada "
                   f"en este tipo.")
    if others:
        out.append("También fue evaluado en " + join_es([f"{ot.type_name.lower()} ({pct(orow['final'])})"
                                                        for ot, orow in others]) + ".")
    return out
