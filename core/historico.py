"""Comparación de varios periodos a la vez (histórico)."""
from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analisis import CAT_HEX, CATS, PROC, PROC_ORDER, cat_of, fnum, join_es, pct, pkey, pp, share


def guess_label(filename: str) -> str:
    """Adivina el periodo desde el nombre del archivo: 'Evaluacion-2024.xlsx' -> '2024', '2025-1' -> '2025-1'.
    Los nombres con fecha y hora de la plataforma (202609281621) no se adivinan."""
    m = re.search(r"(?<!\d)(20\d{2})(?:\s*[-_ ]\s*([12]))?(?!\d)", filename.rsplit(".", 1)[0])
    if not m:
        return ""
    return m.group(1) + (f"-{m.group(2)}" if m.group(2) else "")


def period_summary(timeline) -> list[dict]:
    out = []
    for label, a in timeline:
        good = a["counts"]["EXCELENTE"] + a["counts"]["MUY BUENO"]
        out.append({"Periodo": label, "Evaluaciones": a["total"], "Proveedores": a["uniq"],
                    "Promedio": a["avg"], "% excelente o muy bueno": share(good, a["total"]),
                    "Bajo 80%": sum(r["final"] < 80 for r, _ in a["all"]),
                    **{c: a["counts"][c] for c in CATS}})
    return out


def area_matrix(timeline) -> dict:
    """{area: {periodo: (promedio, n)}} en el orden de PROC_ORDER."""
    m = {}
    for k in PROC_ORDER:
        row = {}
        for label, a in timeline:
            g = next((g for g in a["procs"] if g["k"] == k), None)
            if g:
                row[label] = (g["avg"], g["n"])
        if row:
            m[PROC[k]] = row
    return m


def type_matrix(timeline) -> dict:
    """{codigo: {'tipo', 'proceso', 'valores': {periodo: (promedio, n)}}}"""
    m = {}
    for label, a in timeline:
        for t in a["types"]:
            e = m.setdefault(t.code, {"tipo": t.type_name, "proceso": PROC[t.proc], "valores": {}})
            e["valores"][label] = (t.avg, t.n)
    return m


def provider_matrix(timeline) -> dict:
    """{clave: {'nombre', 'nit', 'valores': {periodo: promedio de sus evaluaciones}, 'detalle': {periodo: [(tipo, final)]}}}"""
    m = {}
    for label, a in timeline:
        for r, t in a["all"]:
            k = pkey(r)
            e = m.setdefault(k, {"nombre": r["name"], "nit": r["nit"], "detalle": {}})
            e["nombre"] = r["name"]  # el nombre más reciente
            e["detalle"].setdefault(label, []).append((t.type_name, r["final"]))
    for e in m.values():
        e["valores"] = {lbl: sum(f for _, f in v) / len(v) for lbl, v in e["detalle"].items()}
    return m


def recurrent_low(timeline, min_periods: int = 2) -> list[dict]:
    """Proveedores con alguna evaluación bajo 80% en al menos `min_periods` periodos."""
    pm = provider_matrix(timeline)
    out = []
    for e in pm.values():
        lows = {lbl: min(f for _, f in v) for lbl, v in e["detalle"].items() if min(f for _, f in v) < 80}
        if len(lows) >= min_periods:
            out.append({"nombre": e["nombre"], "nit": e["nit"], "bajos": lows})
    return sorted(out, key=lambda x: (-len(x["bajos"]), x["nombre"]))


def _trend(vals: list[float]) -> str:
    if len(vals) < 3:
        return ""
    ups = all(b > a + 0.05 for a, b in zip(vals, vals[1:]))
    downs = all(b < a - 0.05 for a, b in zip(vals, vals[1:]))
    return "sube" if ups else "baja" if downs else ""


def insights(timeline) -> list[str]:
    if len(timeline) < 2:
        return []
    labels = [l for l, _ in timeline]
    first, last = labels[0], labels[-1]
    s = period_summary(timeline)
    out = []
    best = max(s, key=lambda x: x["Promedio"])
    worst = min(s, key=lambda x: x["Promedio"])
    out.append(f"Entre {first} y {last} ({len(labels)} periodos) el promedio general pasó de {pct(s[0]['Promedio'])} "
               f"a {pct(s[-1]['Promedio'])} ({pp(s[-1]['Promedio'] - s[0]['Promedio'])}). El periodo con mejor promedio "
               f"fue {best['Periodo']} ({pct(best['Promedio'])}) y el más bajo, {worst['Periodo']} "
               f"({pct(worst['Promedio'])}).")
    out.append(f"La proporción de evaluaciones excelentes o muy buenas pasó de {fnum(s[0]['% excelente o muy bueno'])}% "
               f"a {fnum(s[-1]['% excelente o muy bueno'])}%, y las evaluaciones por debajo de 80% pasaron de "
               f"{s[0]['Bajo 80%']} a {s[-1]['Bajo 80%']}. Se evaluaron {s[0]['Evaluaciones']} evaluaciones en {first} "
               f"y {s[-1]['Evaluaciones']} en {last}.")
    am = area_matrix(timeline)
    deltas = [(area, v[last][0] - v[first][0]) for area, v in am.items() if first in v and last in v]
    if deltas:
        up = max(deltas, key=lambda x: x[1])
        dn = min(deltas, key=lambda x: x[1])
        txt = f"Por área, la que más mejoró entre {first} y {last} fue {up[0].lower()} ({pp(up[1])})"
        txt += f" y la que más bajó, {dn[0].lower()} ({pp(dn[1])})." if dn[1] < -0.05 else "; ninguna área bajó."
        out.append(txt)
    tm = type_matrix(timeline)
    falling = [e["tipo"] for e in tm.values()
               if _trend([e["valores"][l][0] for l in labels if l in e["valores"]]) == "baja"
               and len(e["valores"]) >= 3]
    rising = [e["tipo"] for e in tm.values()
              if _trend([e["valores"][l][0] for l in labels if l in e["valores"]]) == "sube"
              and len(e["valores"]) >= 3]
    if falling:
        out.append("Tipos de evaluación que bajan periodo tras periodo: " + join_es([f.lower() for f in falling[:6]]) + ".")
    if rising:
        out.append("Tipos de evaluación que mejoran periodo tras periodo: " + join_es([f.lower() for f in rising[:6]]) + ".")
    rl = recurrent_low(timeline)
    if rl:
        names = [f"{x['nombre']} ({len(x['bajos'])} periodos)" for x in rl[:6]]
        out.append(f"{len(rl)} {'proveedor quedó' if len(rl) == 1 else 'proveedores quedaron'} por debajo de 80% en dos o "
                   f"más periodos: {join_es(names)}" + (f" y {len(rl) - 6} más." if len(rl) > 6 else "."))
    else:
        out.append("Ningún proveedor quedó por debajo de 80% en más de un periodo.")
    pm = provider_matrix(timeline)
    always = sum(1 for e in pm.values() if all(l in e["valores"] for l in labels))
    out.append(f"{always} proveedores fueron evaluados en todos los periodos comparados.")
    return out


_HDR_FILL = PatternFill("solid", fgColor="1F3B5C")


def _header(ws, values):
    ws.append(values)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[ws.max_row].height = 30
    ws.freeze_panes = "B2"


def _color_cells(ws, first_col, n):
    for row in ws.iter_rows(min_row=2, min_col=first_col, max_col=first_col + n - 1):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = "0.00"
                c.font = Font(bold=True, color=CAT_HEX[cat_of(c.value)][1:])


def add_period_sheets(wb, timeline, base: str | None = None):
    """Agrega al libro las hojas Área x periodo, Tipo x periodo y Proveedor x periodo.
    El cambio se calcula del periodo `base` (por defecto el primero) al último."""
    labels = [l for l, _ in timeline]
    b, last = (base or labels[0]), labels[-1]
    chg = f"Cambio {b}→{last} (pp)"

    ws = wb.create_sheet("Área x periodo")
    _header(ws, ["Área", *labels, chg])
    for area, v in area_matrix(timeline).items():
        ws.append([area, *[v[l][0] if l in v else None for l in labels],
                   v[last][0] - v[b][0] if b in v and last in v else None])
    _color_cells(ws, 2, len(labels))
    ws.column_dimensions["A"].width = 26

    ws = wb.create_sheet("Tipo x periodo")
    _header(ws, ["Proceso", "Código", "Tipo de proveedor", *labels, chg])
    for code, e in sorted(type_matrix(timeline).items(), key=lambda kv: (kv[1]["proceso"], kv[1]["tipo"])):
        v = e["valores"]
        ws.append([e["proceso"], code, e["tipo"], *[v[l][0] if l in v else None for l in labels],
                   v[last][0] - v[b][0] if b in v and last in v else None])
    _color_cells(ws, 4, len(labels))
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["C"].width = 45

    ws = wb.create_sheet("Proveedor x periodo")
    _header(ws, ["Proveedor", "NIT", *labels, "Periodos evaluado", "Periodos bajo 80%"])
    for e in sorted(provider_matrix(timeline).values(), key=lambda e: e["nombre"]):
        v = e["valores"]
        lows = sum(1 for d in e["detalle"].values() if min(f for _, f in d) < 80)
        ws.append([e["nombre"], e["nit"], *[v.get(l) for l in labels], len(v), lows])
    _color_cells(ws, 3, len(labels))
    ws.column_dimensions["A"].width = 50
    ws.column_dimensions["B"].width = 18


def build_hist_xlsx(timeline) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen por periodo"
    s = period_summary(timeline)
    _header(ws, list(s[0].keys()))
    for r in s:
        ws.append(list(r.values()))
    for row in ws.iter_rows(min_row=2):
        row[3].number_format = "0.00"
        row[4].number_format = '0.0"%"'
    ws.column_dimensions["A"].width = 14
    add_period_sheets(wb, timeline)
    ws = wb.create_sheet("Análisis")
    ws.column_dimensions["A"].width = 140
    ws.append(["Análisis del histórico"])
    ws["A1"].font = Font(bold=True, size=13)
    for t in insights(timeline):
        ws.append([t])
        ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
